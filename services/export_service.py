"""Single export service for DataFrame tables.

Task 3 audit
------------
Existing export paths before this service:

1. ``dataframe_image`` + Chrome in performance pages.
   This is fragile in deployment because it depends on a browser executable.

2. Raw ``df.to_excel()`` in ``utils.helpers.to_excel`` and some pages.
   This loses the conditional formatting shown on screen.

3. Hand-drawn matplotlib PNGs plus grouped ZIP exports in ``variations_hvc.py``
   and ``listing_pos_oos.py``.
   This is the reference approach: no browser dependency, deterministic output,
   and pagination to avoid huge canvases.

This module generalizes path 3 and keeps CSV, Excel, PNG and grouped ZIP exports
behind one API. Existing pages are not rewired in task 3; they will consume this
service progressively in later tasks.
"""

from __future__ import annotations

import io
import math
import re
import zipfile
from collections.abc import Callable, Iterable, Mapping
from typing import Any, Optional

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.patches import Rectangle
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

StyleDict = dict[str, Any]
StyleRule = Callable[[Any, pd.Series, Any], Optional[StyleDict]]

DEFAULT_STYLE: StyleDict = {
    "header_fill": "#1f2937",
    "subheader_fill": "#374151",
    "header_font_color": "#ffffff",
    "row_even_fill": "#ffffff",
    "row_odd_fill": "#f3f4f6",
    "border_color": "#d1d5db",
    "font_color": "#111827",
    "total_fill": "#e5e7eb",
    "total_font_color": "#111827",
    "font_size": 8.0,
    "header_font_size": 8.5,
    "max_rows_per_image": 40,
    "max_image_height_px": 6000,
    "dpi": 200,
    "sheet_name": "Analyse",
}


def to_csv(df: pd.DataFrame, encoding: str = "utf-8-sig") -> bytes:
    """Return a CSV export as bytes."""
    return df.to_csv(index=False).encode(encoding)


def to_excel(
    df: pd.DataFrame,
    styles: Optional[StyleDict] = None,
    sheet_name: str = "Analyse",
) -> bytes:
    """Return an XLSX export with real cell formatting.

    ``styles`` is the same style contract used by ``to_image``:
    - ``column_styles``: ``{column: style_dict}``
    - ``row_styles``: callable ``(row) -> style_dict``
    - ``rules``: list of callables or dict rules
    - ``formats``: ``{column: excel_number_format}``
    - ``total_row_match``: callable ``(row) -> bool``
    """
    style = _merge_style(styles)
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = _safe_sheet_name(sheet_name or style["sheet_name"])

    _write_excel_table(ws, df, style)
    wb.save(output)
    return output.getvalue()


def to_image(
    df: pd.DataFrame,
    styles: Optional[StyleDict] = None,
    title: str = "",
    paginate: bool = False,
) -> bytes:
    """Return a PNG rendering of a DataFrame table.

    By default, large tables are reduced to the first configured page with a
    title note. Use ``to_image_pages`` or ``to_grouped_zip(..., export='image')``
    for exhaustive paginated exports.
    """
    pages = to_image_pages(df, styles=styles, title=title)
    if paginate or len(pages) == 1:
        return pages[0]

    style = _merge_style(styles)
    limit = int(style["max_rows_per_image"])
    page_df = df.head(limit)
    suffix = f" - first {limit} rows out of {len(df)}; use ZIP for full export"
    return _render_image_page(page_df, style, f"{title}{suffix}" if title else suffix.strip(" -"))


def to_image_pages(
    df: pd.DataFrame,
    styles: Optional[StyleDict] = None,
    title: str = "",
) -> list[bytes]:
    """Return all PNG pages for a DataFrame table."""
    style = _merge_style(styles)
    max_rows = int(style["max_rows_per_image"])
    if df.empty or len(df) <= max_rows:
        return [_render_image_page(df, style, title)]

    pages: list[bytes] = []
    page_count = math.ceil(len(df) / max_rows)
    for page_index in range(page_count):
        start = page_index * max_rows
        chunk = df.iloc[start : start + max_rows]
        page_title = f"{title} (page {page_index + 1}/{page_count})" if title else f"page {page_index + 1}/{page_count}"
        pages.append(_render_image_page(chunk, style, page_title))
    return pages


def to_zip(files: Mapping[str, bytes]) -> bytes:
    """Return a ZIP built from an in-memory mapping of filename to bytes."""
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files.items():
            if data is None:
                continue
            zf.writestr(_safe_archive_name(name), data)
    return output.getvalue()


def to_grouped_zip(
    df: pd.DataFrame,
    group_col: Any,
    export: str = "image",
    styles: Optional[StyleDict] = None,
    filename_prefix: str = "export",
) -> bytes:
    """Export one file per group and return them as a ZIP.

    ``export`` can be ``"image"``, ``"excel"`` or ``"csv"``. Image exports are
    automatically paginated inside each group.
    """
    if df.empty or group_col not in df.columns:
        return to_zip({})

    files: dict[str, bytes] = {}
    for group_value in sorted(df[group_col].dropna().unique().tolist(), key=lambda value: str(value)):
        group_df = df[df[group_col] == group_value].copy()
        clean_group = _safe_file_part(group_value)

        if export == "csv":
            files[f"{filename_prefix}_{clean_group}.csv"] = to_csv(group_df)
        elif export == "excel":
            files[f"{filename_prefix}_{clean_group}.xlsx"] = to_excel(
                group_df,
                styles=styles,
                sheet_name=str(group_value),
            )
        elif export == "image":
            pages = to_image_pages(group_df, styles=styles, title=str(group_value))
            for idx, png in enumerate(pages, start=1):
                suffix = f"_p{idx}" if len(pages) > 1 else ""
                files[f"{filename_prefix}_{clean_group}{suffix}.png"] = png
        else:
            raise ValueError("export must be one of: image, excel, csv")

    return to_zip(files)


def _write_excel_table(ws, df: pd.DataFrame, style: StyleDict) -> None:
    columns = list(df.columns)
    header_rows = _header_depth(columns)

    _write_excel_headers(ws, columns, style, header_rows)
    ws.freeze_panes = ws.cell(row=header_rows + 1, column=1)

    formats = style.get("formats", {})
    for row_offset, (_, row) in enumerate(df.iterrows(), start=header_rows + 1):
        is_total = _is_total_row(row, style)
        for col_idx, col in enumerate(columns, start=1):
            value = row[col]
            cell = ws.cell(row=row_offset, column=col_idx, value=_excel_value(value))
            cell.alignment = Alignment(horizontal=_excel_alignment(value))

            resolved = _resolve_cell_style(value, row, col, style, is_total=is_total)
            _apply_openpyxl_style(cell, resolved)

            fmt = _lookup_col_mapping(formats, col)
            if fmt:
                cell.number_format = fmt

    for col_idx, col in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        sample_values = [_display_value(v) for v in df[col].head(100).tolist()] if col in df else []
        width = max([len(_column_label(col)), *[len(v) for v in sample_values], 8]) + 2
        ws.column_dimensions[letter].width = min(width, 42)


def _write_excel_headers(ws, columns: list[Any], style: StyleDict, header_rows: int) -> None:
    header_font = Font(color=_hex(style["header_font_color"]), bold=True)
    header_fill = PatternFill(start_color=_hex(style["header_fill"]), end_color=_hex(style["header_fill"]), fill_type="solid")
    subheader_fill = PatternFill(start_color=_hex(style["subheader_fill"]), end_color=_hex(style["subheader_fill"]), fill_type="solid")

    if header_rows == 1:
        for col_idx, col in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=_column_label(col))
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        return

    tuples = [_as_tuple(col) for col in columns]
    for level in range(header_rows):
        col_idx = 1
        while col_idx <= len(tuples):
            label = tuples[col_idx - 1][level] if level < len(tuples[col_idx - 1]) else ""
            span_end = col_idx
            if level == 0:
                while span_end < len(tuples) and tuples[span_end][level] == label:
                    span_end += 1
            if span_end > col_idx:
                ws.merge_cells(start_row=level + 1, start_column=col_idx, end_row=level + 1, end_column=span_end)
            cell = ws.cell(row=level + 1, column=col_idx, value=_clean_label(label))
            cell.font = header_font
            cell.fill = header_fill if level == 0 else subheader_fill
            cell.alignment = Alignment(horizontal="center")
            col_idx = span_end + 1 if span_end > col_idx else col_idx + 1


def _render_image_page(df: pd.DataFrame, style: StyleDict, title: str = "") -> bytes:
    columns = list(df.columns)
    if df.empty or not columns:
        fig, ax = plt.subplots(figsize=(6, 2), dpi=int(style["dpi"]))
        ax.text(0.5, 0.5, "Aucune donnee a exporter", ha="center", va="center", color=style["font_color"])
        ax.axis("off")
        return _fig_to_png(fig)

    header_rows = _header_depth(columns)
    widths = [_image_col_width(df, col, style) for col in columns]
    total_w = sum(widths)
    n_rows = len(df)

    fig_w = max(8.0, total_w * 0.68)
    fig_h = max(2.2, (n_rows + header_rows) * 0.34 + (0.6 if title else 0.2))
    dpi = int(style["dpi"])
    if fig_h * dpi > int(style["max_image_height_px"]):
        dpi = max(60, int(int(style["max_image_height_px"]) / fig_h))

    fig, ax = plt.subplots(figsize=(fig_w, fig_h), dpi=dpi)
    ax.set_xlim(0, total_w)
    ax.set_ylim(0, n_rows + header_rows)
    ax.invert_yaxis()
    ax.axis("off")

    if title:
        ax.text(
            total_w / 2,
            -0.45,
            title,
            ha="center",
            va="bottom",
            fontsize=10,
            fontweight="bold",
            color=style["font_color"],
        )

    _draw_image_headers(ax, columns, widths, style, header_rows)

    for row_idx, (_, row) in enumerate(df.iterrows()):
        y = header_rows + row_idx
        base_fill = style["row_even_fill"] if row_idx % 2 == 0 else style["row_odd_fill"]
        is_total = _is_total_row(row, style)
        x = 0.0

        for col, width in zip(columns, widths):
            value = row[col]
            resolved = {"fill": base_fill, "font_color": style["font_color"]}
            resolved.update(_resolve_cell_style(value, row, col, style, is_total=is_total))

            rect = Rectangle(
                (x, y),
                width,
                1,
                facecolor=resolved.get("fill", base_fill),
                edgecolor=style["border_color"],
                linewidth=0.5,
            )
            ax.add_patch(rect)

            if str(col) == "Historique":
                symbols = str(value).split()
                if symbols:
                    n_sym = len(symbols)
                    box_w = min(0.22, (width * 0.7) / max(1, n_sym))
                    box_h = 0.45
                    gap = 0.05
                    total_boxes_w = n_sym * box_w + (n_sym - 1) * gap
                    start_x = x + (width - total_boxes_w) / 2.0
                    by = y + (1.0 - box_h) / 2.0
                    for idx_sym, sym in enumerate(symbols):
                        bx = start_x + idx_sym * (box_w + gap)
                        color = "#e74c3c" if "🟥" in sym else ("#2ecc71" if "🟩" in sym else "#d1d5db")
                        patch = Rectangle((bx, by), box_w, box_h, facecolor=color, edgecolor="#ffffff", linewidth=0.6)
                        ax.add_patch(patch)
            else:
                text = _display_value(value, col, style)
                align = resolved.get("align") or _image_alignment(value)
                tx = x + 0.08 if align == "left" else x + width / 2
                text_obj = ax.text(
                    tx,
                    y + 0.5,
                    _fit_text(text, width, fig_w, total_w, float(style["font_size"])),
                    ha=align,
                    va="center",
                    fontsize=float(style["font_size"]),
                    color=resolved.get("font_color", style["font_color"]),
                    fontweight="bold" if resolved.get("bold") else "normal",
                )
                text_obj.set_clip_path(rect)
            x += width

    fig.tight_layout(pad=0.6)
    return _fig_to_png(fig)


def _draw_image_headers(ax, columns: list[Any], widths: list[float], style: StyleDict, header_rows: int) -> None:
    if header_rows == 1:
        x = 0.0
        for col, width in zip(columns, widths):
            rect = Rectangle((x, 0), width, 1, facecolor=style["header_fill"], edgecolor=style["border_color"], linewidth=0.6)
            ax.add_patch(rect)
            text = ax.text(
                x + width / 2,
                0.5,
                _clean_label(_column_label(col)),
                ha="center",
                va="center",
                fontsize=float(style["header_font_size"]),
                fontweight="bold",
                color=style["header_font_color"],
            )
            text.set_clip_path(rect)
            x += width
        return

    tuples = [_as_tuple(col) for col in columns]
    for level in range(header_rows):
        y = level
        x = 0.0
        idx = 0
        while idx < len(columns):
            label = tuples[idx][level] if level < len(tuples[idx]) else ""
            span_width = widths[idx]
            span_end = idx
            if level == 0:
                while span_end + 1 < len(columns) and tuples[span_end + 1][level] == label:
                    span_end += 1
                    span_width += widths[span_end]

            rect = Rectangle(
                (x, y),
                span_width,
                1,
                facecolor=style["header_fill"] if level == 0 else style["subheader_fill"],
                edgecolor=style["border_color"],
                linewidth=0.6,
            )
            ax.add_patch(rect)
            text = ax.text(
                x + span_width / 2,
                y + 0.5,
                _clean_label(str(label)),
                ha="center",
                va="center",
                fontsize=float(style["header_font_size"]),
                fontweight="bold",
                color=style["header_font_color"],
            )
            text.set_clip_path(rect)
            x += span_width
            idx = span_end + 1


def _resolve_cell_style(value: Any, row: pd.Series, col: Any, style: StyleDict, is_total: bool = False) -> StyleDict:
    resolved: StyleDict = {}

    col_style = _lookup_col_mapping(style.get("column_styles", {}), col)
    if col_style:
        resolved.update(col_style)

    row_style = style.get("row_styles")
    if callable(row_style):
        resolved.update(row_style(row) or {})

    for rule in style.get("rules", []):
        if callable(rule):
            resolved.update(rule(value, row, col) or {})
        elif isinstance(rule, Mapping) and _rule_matches(rule, value, row, col):
            resolved.update(rule.get("style", {}))

    cell_styles = style.get("cell_styles", {})
    cell_style = _lookup_col_mapping(cell_styles, col)
    if isinstance(cell_style, Mapping):
        resolved.update(cell_style.get(row.name, {}) if row.name in cell_style else {})

    if is_total:
        resolved.setdefault("fill", style["total_fill"])
        resolved.setdefault("font_color", style["total_font_color"])
        resolved["bold"] = True

    return resolved


def _rule_matches(rule: Mapping[str, Any], value: Any, row: pd.Series, col: Any) -> bool:
    columns = rule.get("columns")
    if columns is not None and not _col_in(columns, col):
        return False

    when = rule.get("when")
    if callable(when):
        return bool(when(value, row, col))

    op = rule.get("op")
    target = rule.get("value")
    try:
        if op == "<":
            return value < target
        if op == "<=":
            return value <= target
        if op == ">":
            return value > target
        if op == ">=":
            return value >= target
        if op == "==":
            return value == target
        if op == "!=":
            return value != target
        if op == "between":
            low, high = target
            return low <= value <= high
    except Exception:
        return False
    return False


def _merge_style(styles: Optional[StyleDict]) -> StyleDict:
    merged = DEFAULT_STYLE.copy()
    if styles:
        merged.update(styles)
    return merged


def _apply_openpyxl_style(cell, style: StyleDict) -> None:
    fill = style.get("fill")
    if fill:
        cell.fill = PatternFill(start_color=_hex(fill), end_color=_hex(fill), fill_type="solid")

    font_kwargs: dict[str, Any] = {}
    if style.get("font_color"):
        font_kwargs["color"] = _hex(style["font_color"])
    if style.get("bold"):
        font_kwargs["bold"] = True
    if font_kwargs:
        cell.font = Font(**font_kwargs)

    if style.get("align"):
        cell.alignment = Alignment(horizontal=style["align"])


def _header_depth(columns: Iterable[Any]) -> int:
    return max((len(_as_tuple(col)) for col in columns), default=1)


def _as_tuple(col: Any) -> tuple[Any, ...]:
    return col if isinstance(col, tuple) else (col,)


def _column_label(col: Any) -> str:
    if isinstance(col, tuple):
        return " | ".join(str(part) for part in col if str(part) != "")
    return str(col)


def _clean_label(label: Any) -> str:
    text = str(label)
    text = re.sub(r"[\U0001F300-\U0001FAFF\U00002600-\U000027BF\U0001F000-\U0001F0FF]+", "", text)
    return text.strip()


def _display_value(value: Any, col: Any = None, style: Optional[StyleDict] = None) -> str:
    if pd.isna(value):
        return "-"

    formatter = _lookup_col_mapping((style or {}).get("display_formats", {}), col) if style else None
    if callable(formatter):
        return str(formatter(value))

    if isinstance(value, float):
        return f"{value:,.2f}" if not value.is_integer() else f"{value:,.0f}"
    if isinstance(value, int):
        return f"{value:,}"
    return str(value)


def _excel_value(value: Any) -> Any:
    if pd.isna(value):
        return None
    if hasattr(value, "item"):
        try:
            return value.item()
        except Exception:
            return value
    return value


def _image_col_width(df: pd.DataFrame, col: Any, style: StyleDict) -> float:
    configured = _lookup_col_mapping(style.get("column_widths", {}), col)
    if configured:
        return float(configured)
    header_len = len(_clean_label(_column_label(col)))
    sample = df[col].head(50).map(lambda value: len(_display_value(value, col, style))).max() if col in df else 8
    max_len = max(header_len, int(sample) if pd.notna(sample) else 8)
    return min(max(0.9, max_len * 0.11), 3.8)


def _fit_text(text: str, cell_width: float, fig_w: float, total_w: float, font_size: float) -> str:
    points_per_unit = (fig_w / total_w) * 72
    max_chars = max(3, int((cell_width * points_per_unit - 6) / (font_size * 0.56)))
    if len(text) <= max_chars:
        return text
    if max_chars <= 1:
        return text[:max_chars]
    return text[: max_chars - 1].rstrip() + "..."


def _image_alignment(value: Any) -> str:
    return "center" if isinstance(value, (int, float)) and not isinstance(value, bool) else "left"


def _excel_alignment(value: Any) -> str:
    return "center" if isinstance(value, (int, float)) and not isinstance(value, bool) else "left"


def _lookup_col_mapping(mapping: Mapping[Any, Any], col: Any) -> Any:
    if not mapping:
        return None
    if col in mapping:
        return mapping[col]
    label = _column_label(col)
    if label in mapping:
        return mapping[label]
    if isinstance(col, tuple) and col[-1] in mapping:
        return mapping[col[-1]]
    return None


def _col_in(columns: Iterable[Any], col: Any) -> bool:
    if col in columns:
        return True
    label = _column_label(col)
    return label in columns or (isinstance(col, tuple) and col[-1] in columns)


def _is_total_row(row: pd.Series, style: StyleDict) -> bool:
    matcher = style.get("total_row_match")
    if callable(matcher):
        return bool(matcher(row))
    return any(str(value).strip().lower() == "total" for value in row.tolist()[:2])


def _safe_sheet_name(value: str) -> str:
    clean = re.sub(r"[\\/*?:\[\]]", "_", str(value)).strip() or "Analyse"
    return clean[:31]


def _safe_file_part(value: Any, max_length: int = 80) -> str:
    clean = re.sub(r"[^\w\-. ]", "_", str(value)).strip().replace(" ", "_")
    return (clean or "groupe")[:max_length]


def _safe_archive_name(value: str) -> str:
    parts = [part for part in str(value).replace("\\", "/").split("/") if part not in ("", ".", "..")]
    return "/".join(parts) if parts else "export.bin"


def _hex(color: str) -> str:
    color = str(color).strip()
    if color.startswith("#"):
        color = color[1:]
    if len(color) == 6:
        return color.upper()
    if len(color) == 8:
        return color[-6:].upper()
    return color.upper()


def _fig_to_png(fig) -> bytes:
    output = io.BytesIO()
    fig.savefig(output, format="png", bbox_inches="tight")
    plt.close(fig)
    return output.getvalue()
